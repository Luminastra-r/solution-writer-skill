#!/usr/bin/env python3
"""Reproducible offline orchestrator sample with public data, no model/API expense."""
import argparse
import hashlib
import json
import sys
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "tests"))
from sample_solution import fixture
from orchestrate_solution import orchestrate, build_arg_parser
from docx import Document
from docx.oxml.ns import qn


class OfflineLLM:
    def __init__(self, **kwargs):
        _, self.blueprint, self.bodies, _ = fixture()

    def generate(self, **kwargs):
        kwargs["state_store"].increment_llm_calls()
        key = kwargs["step_key"]
        if key == "blueprint":
            return json.dumps(self.blueprint, ensure_ascii=False)
        if key.startswith("chapter_"):
            return self.bodies[key[len("chapter_"):]]
        raise RuntimeError(f"Unexpected LLM call: {key}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", default=str(REPO/"artifacts"/"visualization-regression"))
    parser.add_argument("--engine", choices=["auto", "svg", "mmdc"], default="svg")
    parser.add_argument("--mmdc-bin")
    parser.add_argument("--gradient-mode", choices=["controlled", "solid"], default="solid")
    parser.add_argument("--check-only", action="store_true", help="Check existing outputs without regeneration")
    args = parser.parse_args()
    out = Path(args.output_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    request, _, _, _ = fixture()
    request["visualization"].update(engine=args.engine, mmdc_bin=args.mmdc_bin, gradient_mode=args.gradient_mode)
    request_path = out/"request.json"
    if not args.check_only:
        request_path.write_text(json.dumps(request, ensure_ascii=False, indent=2), encoding="utf-8")
        cli_args = build_arg_parser().parse_args(["--input-json", str(request_path), "--output-dir", str(out),
            "--model", "offline-fixture", "--run-mode", "fast", "--research-mode", "off"])
        with patch("orchestrate_solution.OpenAILLM", OfflineLLM):
            result = orchestrate(cli_args)
        if result:
            return result
    state = json.loads((out/"run_state.json").read_text(encoding="utf-8"))
    manifest = json.loads((out/"diagrams"/"diagram_manifest.json").read_text(encoding="utf-8"))
    doc = Document(out/"sample_solution.docx")
    with ZipFile(out/"sample_solution.docx") as archive:
        image_hashes = {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith("word/media/")}
    inserted = [i for i in manifest["items"] if i.get("insertion_status") == "inserted"]
    png_hashes = {i["diagram_id"]: hashlib.sha256(Path(i["image_path"]).read_bytes()).hexdigest() for i in inserted}
    spacings = [p._p.pPr.find(qn("w:spacing")) for t in doc.tables if len(t.rows) > 1
                for r in t.rows for c in r.cells for p in c.paragraphs]
    checks = {
        "mode": "actual orchestrator, deterministic offline LLM fixture",
        "llm_calls": state["llm_call_count"],
        "figure_count": sum(i.get("insertion_status") == "inserted" for i in manifest["items"]),
        "docx_captions": sum(p.text.startswith("图") for p in doc.paragraphs),
        "ordinary_table_paragraphs": len(spacings),
        "table_multiple_1_2": all(s.get(qn("w:line")) == "288" and s.get(qn("w:lineRule")) == "auto" for s in spacings),
        "body_spacing_1_75": next(p for p in doc.paragraphs if p.text.startswith("工信部发布")).paragraph_format.line_spacing == 1.75,
        "gradient_presentation_count": sum(i.get("gradient_mode") == "controlled" and i["type"] in ("architecture", "business_model") for i in inserted),
        "gradient_data_count": sum(i.get("gradient_mode") == "controlled" and i["type"] == "data_chart" for i in inserted),
        "gradient_fallbacks": [i["diagram_id"] for i in inserted if i.get("gradient_mode") == "solid_fallback"],
        "docx_embedded_png_byte_identical": bool(png_hashes) and all(h in image_hashes for h in png_hashes.values()),
        "png_sha256": png_hashes,
        "docx_page_render": "failed: bundled renderer cannot find LibreOffice soffice.exe" if (out/"docx-render.log").is_file() and "soffice.exe was not found" in (out/"docx-render.log").read_text(encoding="utf-8") else "not run; requires an available bundled LibreOffice renderer",
        "live_llm_generation": "not run; offline fixture exercises orchestration without an external model",
    }
    (out/"checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    return 0 if checks["figure_count"] == checks["docx_captions"] == 5 and checks["table_multiple_1_2"] else 2


if __name__ == "__main__":
    sys.exit(main())
