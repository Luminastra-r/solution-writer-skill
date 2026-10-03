"""Fixed public-data comparisons; run before/after into separate retained folders."""
import argparse
import hashlib
import json
import sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"tests"))
from sample_solution import fixture, graph
from solution_skill.visualization.pipeline import render_tasks
from generate_docx import create_solution_docx
from PIL import Image
from docx import Document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    out = Path(args.output_dir).resolve()
    request, _, _, original = fixture()
    request["visualization"].update(theme="sage", gradient_mode="solid")
    source = "https://www.miit.gov.cn/jgsj/yxj/xxfb/art/2024/art_3cb679c2662d4127af3cc857d7dbff8e.html"
    request["visualization_data"].append({"id":"annual", "labels":["2023年","2024年"],
        "series":[{"name":"软件产品", "values":[29030,30417]}, {"name":"信息技术服务", "values":[81226,92190]}],
        "unit":"亿元", "source":{"kind":"verified","reference":source+"；"+request["visualization_data"][0]["source"]["reference"]}})
    tasks = [original[1], graph("capabilities", "architecture", "ch01", "", "四模块业务能力", ["业务咨询","事项办理","数据治理","制度规范"], edges=[]),
             original[0], dict(original[0], id="annual_line", title="全国软件业两类收入趋势", chart="line", dataset_id="annual")]
    tasks[1]["theme"] = "tech"
    out.mkdir(parents=True, exist_ok=True)
    (out/"fixture.json").write_text(json.dumps({"request":request,"diagrams":tasks},ensure_ascii=False,indent=2),encoding="utf-8")
    manifest = render_tasks(tasks, out/"diagrams", request)
    md = "# 图表视觉对比样本\n\n以下业务结构为方案设计，行业数值不是客户经营成效。2023与2024年均为当年公开口径，不推导可比增速。\n\n"
    for item in manifest["items"]:
        if item["status"] != "ok":
            raise RuntimeError(item)
        md += f'## {item["title"]}\n\n![{item["title"]}]({Path(item["image_path"]).as_posix()})\n\n'
        md += '<!-- solution-source:'+json.dumps(item['source']['reference'],ensure_ascii=False)+' -->\n\n'
    md += '| 检查项目 | 说明 |\n| --- | --- |\n| 表格行距 | 保持1.2倍 |\n'
    (out/"solution.md").write_text(md,encoding="utf-8")
    create_solution_docx(md,str(out/"sample.docx"),"图表视觉对比","验收样本",str(out))
    # Independent diagnostic: deliberately strong gradient, never a production theme.
    probe = out/"gradient_probe.svg"
    probe.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="680" height="180" viewBox="0 0 680 180"><defs><linearGradient id="probe" gradientUnits="userSpaceOnUse" x1="20" y1="0" x2="660" y2="0"><stop offset="0" stop-color="#5B6CFF"/><stop offset="1" stop-color="#7B61FF"/></linearGradient></defs><rect width="680" height="180" fill="white"/><rect x="20" y="20" width="640" height="140" fill="url(#probe)"/></svg>',encoding="utf-8")
    from svg_to_png import convert_with_cairosvg
    convert_with_cairosvg(probe,out/"gradient_probe.png",1800)
    with Image.open(out/"gradient_probe.png") as im:
        samples=[im.convert("RGB").getpixel((round(x*im.width/680),round(90*im.width/680))) for x in (60,200,340,480,620)]
    create_solution_docx(f'## 渐变链路诊断\n\n![渐变诊断]({(out/"gradient_probe.png").as_posix()})',str(out/"gradient_probe.docx"),"渐变链路诊断","验收样本",str(out))
    with ZipFile(out/"gradient_probe.docx") as z:
        digest=hashlib.sha256((out/"gradient_probe.png").read_bytes()).hexdigest()
        identical=any(hashlib.sha256(z.read(n)).hexdigest()==digest for n in z.namelist() if n.startswith('word/media/'))
    checks={"svg_gradient_definition_and_reference":'url(#probe)' in probe.read_text(),"gradient_units":"userSpaceOnUse",
            "png_internal_samples_rgb":samples,"png_gradient_preserved":len(set(samples))==5,
            "docx_embedded_png_identical":identical,"final_docx_page_render":"NOT VERIFIED; run packaged renderer separately",
            "gradient_full_chain_accepted":False,"production_policy":"solid by default until final page verification",
            "figure_count":len(manifest["items"])}
    doc = Document(out/"sample.docx")
    checks['table_spacing_1_2'] = all(p.paragraph_format.line_spacing == 1.2 for t in doc.tables if len(t.rows)>1 for row in t.rows for cell in row.cells for p in cell.paragraphs)
    checks['inserted_sizes_and_estimated_min_font_pt'] = []
    for shape, item, minimum in zip(doc.inline_shapes,manifest['items'],(16*640/680,17*640/680,11,11)):
        viewbox=ET.parse(item['svg_path']).getroot().get('viewBox').split()
        font_pt=minimum*shape.width.pt/float(viewbox[2])
        checks['inserted_sizes_and_estimated_min_font_pt'].append({'id':item['diagram_id'],'width_cm':shape.width.cm,'height_cm':shape.height.cm,'min_font_pt_estimate':round(font_pt,2)})
        with Image.open(item['image_path']) as im:
            im.resize((round(shape.width.inches*144),round(shape.height.inches*144))).save(out/f'{item["diagram_id"]}-at-144dpi.png')
    with ZipFile(out/'sample.docx') as z:
        hashes={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    checks['all_four_embedded_png_identical']=all(hashlib.sha256(Path(i['image_path']).read_bytes()).hexdigest() in hashes for i in manifest['items'])
    baseline=out.parent/'before'/'fixture.json'
    if baseline.exists():
        checks['identical_before_after_content']=json.loads(baseline.read_text(encoding='utf-8'))==json.loads((out/'fixture.json').read_text(encoding='utf-8'))
    if (out/'gradient-render.log').exists():
        checks['final_docx_page_render']='FAILED: LibreOffice soffice.exe unavailable; see gradient-render.log and docx-render.log'
    (out/"checks.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(checks,ensure_ascii=False,indent=2))

if __name__ == "__main__":
    main()
