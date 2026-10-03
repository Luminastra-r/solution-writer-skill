from copy import deepcopy
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import patch

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

from inject_diagrams import inject_items, locate
from render_diagrams import build_mermaid
from solution_skill.visualization.planning import validate_task, extract_tasks, options
from solution_skill.visualization.pipeline import render_tasks, visualize_solution
from solution_skill.writing.blueprint import flatten_chapters
from solution_skill.writing.markdown_builder import build_solution_markdown
from generate_docx import create_solution_docx, set_table_line_spacing
from sample_solution import fixture, graph


def test_four_types_and_same_section_two_figures(tmp_path):
    req, bp, bodies, tasks = fixture()
    bodies = {cid: extract_tasks(text)[0] for cid, text in bodies.items()}
    md = build_solution_markdown(bp, flatten_chapters(bp), bodies)
    rendered, manifest = visualize_solution(md, tasks, req, {}, tmp_path)
    assert len(manifest["items"]) == 5
    assert all(i["status"] == "ok" and i["insertion_status"] == "inserted" for i in manifest["items"])
    assert {i["type"] for i in manifest["items"]} == {"flowchart", "architecture", "data_chart", "business_model"}
    same = rendered.split("<!-- solution-section:ch03_s1 -->")[1].split("<!-- solution-section:ch03_s2 -->")[0]
    assert same.count("![") == 2
    assert "![" not in rendered.split("<!-- solution-chapter:ch02 -->")[1].split("<!-- solution-chapter:ch03 -->")[0]
    for item in manifest["items"]:
        assert ET.parse(item["svg_path"]).getroot().tag.endswith("svg")
        with Image.open(item["image_path"]) as image:
            assert image.width >= 1200
    # Repeated render/injection must replace blocks including source and caption.
    repeated, _ = visualize_solution(rendered, tasks, req, {}, tmp_path)
    assert repeated.count("![") == rendered.count("![") == 5
    assert repeated.count("solution-source:") == 5
    create_solution_docx(repeated, str(tmp_path/"sample.docx"), bp["title"], "测试园区", str(tmp_path))
    doc = Document(tmp_path/"sample.docx")
    assert sum(p.text.startswith("图") for p in doc.paragraphs) == 5
    assert not any("solution-" in p.text or "visualizations" in p.text for p in doc.paragraphs)
    ordinary = [t for t in doc.tables if len(t.rows) > 1]
    assert ordinary
    for table in ordinary:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    assert p.paragraph_format.line_spacing == 1.2
                    assert p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
                    spacing = p._p.pPr.find(qn("w:spacing"))
                    assert spacing.get(qn("w:line")) == "288"
                    assert spacing.get(qn("w:lineRule")) == "auto"
    assert next(p for p in doc.paragraphs if p.text.startswith("工信部发布")).paragraph_format.line_spacing == 1.75


@pytest.mark.parametrize("chart", ["bar", "line", "pie", "scatter"])
def test_data_subtypes(chart, tmp_path):
    req, _, _, tasks = fixture()
    d = dict(tasks[0], chart=chart)
    m = render_tasks([d], tmp_path, req)
    assert m["items"][0]["status"] == "ok", m


@pytest.mark.parametrize("change", ["missing", "assumed", "nan", "mismatch", "pie_negative"])
def test_no_unreliable_data(change):
    req, _, _, tasks = fixture()
    d = deepcopy(tasks[0])
    if change == "missing":
        d["dataset_id"] = "invented"
    elif change == "assumed":
        req["visualization_data"][0]["source"]["kind"] = "assumed"
    elif change == "nan":
        req["visualization_data"][0]["series"][0]["values"][0] = float("nan")
    elif change == "mismatch":
        req["visualization_data"][0]["series"][0]["values"].pop()
    else:
        d["chart"] = "pie"
        req["visualization_data"][0]["series"][0]["values"][0] = -1
    with pytest.raises(ValueError):
        validate_task(d, req)


def test_numeric_quotes_verified_and_user():
    d = {"id": "revenue", "type": "data_chart", "title": "资料收入", "points": [
        {"label": "软件产品", "value": 30417, "evidence": "软件产品收入30417亿元"}],
         "source": {"kind": "user_provided", "reference": "用户输入"}}
    assert validate_task(d, {"raw_input": "已知软件产品收入30417亿元。"})["series"][0]["values"] == [30417]
    with pytest.raises(ValueError):
        validate_task(d, {"raw_input": "暂无数据"})
    d["source"] = {"kind": "verified", "reference": "https://example.org/official"}
    pack = {"web_items": [{"link": "https://example.org/official", "snippet": "软件产品收入30417亿元"}]}
    assert validate_task(d, {}, pack)["labels"] == ["软件产品"]
    d["points"][0]["value"] = 88888
    with pytest.raises(ValueError):
        validate_task(d, {}, pack)
    d["source"] = {"kind": "user_provided", "reference": "用户输入"}
    d["points"] = [{"label": "事项A", "value": 20, "evidence": "事项A办理10件，事项B办理20件"}]
    with pytest.raises(ValueError, match="pairing"):
        validate_task(d, {"raw_input": "事项A办理10件，事项B办理20件"})


def test_safe_specs_and_no_silent_truncation():
    d = graph("dense", "architecture", "ch01", "ch01_s1", "结构", [f"节点{i}" for i in range(19)])
    assert "节点18" in build_mermaid(d)
    with pytest.raises(ValueError, match="split"):
        validate_task(d, {})
    d["id"] = "../../escape"
    with pytest.raises(ValueError):
        validate_task(d, {})
    d = graph("good", "flowchart", "ch01", "ch01_s1", "结构", ["前置", "后置"])
    d["edges"][0]["to"] = "unknown"
    with pytest.raises(ValueError):
        build_mermaid(d)
    d["edges"] = []
    src = build_mermaid(d)
    assert src.startswith("%%{init:") and "}%%\nflowchart" in src
    assert "classDef end " not in src and "classDef finish " in src


def test_failure_nonblocking_and_bounded_retry(tmp_path):
    req, bp, bodies, tasks = fixture()
    md = build_solution_markdown(bp, flatten_chapters(bp), {k: extract_tasks(v)[0] for k, v in bodies.items()})
    with patch("solution_skill.visualization.pipeline.render", side_effect=RuntimeError("simulated failure")) as render:
        rendered, manifest = visualize_solution(md, tasks[:1], req, {}, tmp_path)
    assert render.call_count == 2
    assert manifest["items"][0]["status"] == "failed"
    assert "![" not in rendered and "工信部发布" in rendered
    create_solution_docx(rendered, str(tmp_path/"no_figures.docx"), bp["title"], "测试园区", str(tmp_path))
    assert (tmp_path/"no_figures.docx").is_file()


def test_budget_disabled_and_duplicate(tmp_path):
    req, _, _, tasks = fixture()
    req["visualization"]["max_diagrams"] = 1
    m = render_tasks(tasks, tmp_path, req)
    assert sum(i["status"] == "ok" for i in m["items"]) == 1
    req["visualization"]["enabled"] = False
    assert all(i["attempts"] == 0 for i in render_tasks(tasks, tmp_path, req)["items"])
    req["visualization"] = {"engine": "svg"}
    m = render_tasks([tasks[2], dict(tasks[2], id="duplicate")], tmp_path, req)
    assert m["items"][1]["status"] == "skipped"


def test_missing_and_ambiguous_anchor(tmp_path):
    md = "<!-- solution-chapter:ch01 -->\n## 标题\n正文\n正文\n"
    item = {"chapter_id": "ch01", "anchor_text": "正文"}
    with pytest.raises(ValueError, match="ambiguous"):
        locate(md.splitlines(), item)
    with pytest.raises(ValueError):
        locate(md.splitlines(), dict(item, section_id="unknown"))
    assert locate(md.splitlines(), {"chapter_id": "ch01"}) == 4


def test_malformed_metadata_preserves_prose():
    body, tasks, errors = extract_tasks("正文保留\n```visualizations\n{bad json\n```")
    assert body == "正文保留" and not tasks and errors


def test_table_multiple_paragraphs_only_spacing_changes():
    doc = Document()
    p = doc.add_paragraph("正文")
    p.paragraph_format.line_spacing = 1.75
    table = doc.add_table(rows=2, cols=2)
    table.cell(1, 1).add_paragraph("第二段\n换行")
    table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).add_paragraph("嵌套段落")
    before_table = ET.fromstring(table._tbl.xml)
    before_body = p._p.xml
    set_table_line_spacing(table)
    after_table = ET.fromstring(table._tbl.xml)
    # Removing just new line-spacing attrs must make the complete table identical.
    for tree in (before_table, after_table):
        for element in tree.iter():
            if element.text and not element.text.strip():
                element.text = None
            if element.tail and not element.tail.strip():
                element.tail = None
        for ppr in list(tree.iter(qn("w:pPr"))):
            spacing = ppr.find(qn("w:spacing"))
            if spacing is not None:
                spacing.attrib.pop(qn("w:line"), None)
                spacing.attrib.pop(qn("w:lineRule"), None)
                if not spacing.attrib:
                    ppr.remove(spacing)
        for paragraph in tree.iter(qn("w:p")):
            ppr = paragraph.find(qn("w:pPr"))
            if ppr is not None and not len(ppr) and not ppr.attrib:
                paragraph.remove(ppr)
    assert ET.tostring(before_table) == ET.tostring(after_table)
    assert p._p.xml == before_body


def test_split_preserves_all_edges_and_nodes(tmp_path):
    task = graph("large", "flowchart", "ch01", "", "完整流程", [f"节点{i}" for i in range(19)])
    task["edges"].append({"from": "v18", "to": "v0", "label": "反馈"})
    from solution_skill.visualization.planning import split_graph
    parts = split_graph(task, options({"visualization": {"engine": "svg"}}))
    assert len(parts) > 1
    assert {n["id"] for p in parts for n in p["nodes"]} == {n["id"] for n in task["nodes"]}
    assert [e for p in parts for e in p["edges"]] == task["edges"]
    rendered = render_tasks([task], tmp_path, {"visualization": {"engine": "svg"}})
    assert all(i["status"] == "ok" and i["split_from"] == "large" for i in rendered["items"])
    budget = render_tasks([task], tmp_path, {"visualization": {"engine": "svg", "max_diagrams": 1}})
    assert len(budget["items"]) == 1 and budget["items"][0]["status"] == "skipped"


def test_prose_chart_inconsistency_skips_image(tmp_path):
    req, bp, bodies, tasks = fixture()
    md = build_solution_markdown(bp, flatten_chapters(bp), {k: extract_tasks(v)[0] for k, v in bodies.items()})
    md = md.replace("30417", "123")
    result, manifest = visualize_solution(md, tasks[:1], req, {}, tmp_path)
    assert manifest["items"][0]["status"] == "skipped"
    assert "![" not in result


def test_legacy_multi_image_cli(tmp_path):
    import subprocess, sys
    req, _, _, tasks = fixture()
    specs = [{**t, "chapter_id": "", "section_id": "", "chapter": "chapter1"} for t in tasks[1:3]]
    spec = tmp_path/"specs.json"
    spec.write_text(json.dumps({"diagrams": specs}, ensure_ascii=False), encoding="utf-8")
    scripts = Path(__file__).resolve().parents[1]/"scripts"
    result = subprocess.run([sys.executable, str(scripts/"render_diagrams.py"), "--spec-json", str(spec),
        "--out-dir", str(tmp_path/"diagrams"), "--engine", "svg", "--export-svg"], capture_output=True)
    assert result.returncode == 0, result.stderr
    md = tmp_path/"legacy.md"
    md.write_text("## 一、需求理解\n\n正文\n", encoding="utf-8")
    command = [sys.executable, str(scripts/"inject_diagrams.py"), "--input-md", str(md), "--output-md", str(md),
               "--manifest", str(tmp_path/"diagrams"/"diagram_manifest.json")]
    assert subprocess.run(command, capture_output=True).returncode == 0
    assert subprocess.run(command, capture_output=True).returncode == 0
    assert md.read_text(encoding="utf-8").count("![") == 2


def test_missing_font_is_noncritical_failure(tmp_path):
    req, _, _, tasks = fixture()
    with patch("solution_skill.visualization.renderers.cjk_font", side_effect=RuntimeError("font unavailable")):
        manifest = render_tasks(tasks[:1], tmp_path, req)
    assert manifest["items"][0]["status"] == "failed"
    assert manifest["items"][0]["attempts"] == 2


def test_scatter_numeric_x_trusted_dataset(tmp_path):
    req, _, _, tasks = fixture()
    req["visualization_data"][0].update(x_values=[1.2, 2.3, 3.4, 4.5], x_label="已核实横轴")
    task = dict(tasks[0], chart="scatter", numeric_x=True, x_values=[999]*4)
    validated = validate_task(task, req)
    assert validated["x_values"] == [1.2, 2.3, 3.4, 4.5]
    assert render_tasks([task], tmp_path, req)["items"][0]["status"] == "ok"
    req["visualization_data"][0].pop("x_values")
    validated = validate_task(task, req)
    assert not validated.get("numeric_x") and validated["x_values"] == [1, 2, 3, 4]


def test_standalone_docx_cli_table_spacing(tmp_path):
    import subprocess, sys
    scripts = Path(__file__).resolve().parents[1]/"scripts"
    md, output = tmp_path/"input.md", tmp_path/"output.docx"
    md.write_text("## 测试章节\n\n正文保留\n\n| 角色 | 说明 |\n| --- | --- |\n| 运营 | 第一行<br>第二行 |\n", encoding="utf-8")
    result = subprocess.run([sys.executable, str(scripts/"generate_docx.py"), "--input", str(md),
        "--output", str(output), "--project", "项目", "--customer", "客户"], capture_output=True)
    assert result.returncode == 0, result.stderr
    doc = Document(output)
    for table in doc.tables:
        if len(table.rows) > 1:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        assert p.paragraph_format.line_spacing == 1.2
    assert next(p for p in doc.paragraphs if p.text == "正文保留").paragraph_format.line_spacing == 1.75
