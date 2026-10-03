"""Run the actual orchestrator with a deterministic LLM boundary, no remote calls."""
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING

from orchestrate_solution import orchestrate, build_arg_parser
from solution_skill.export.docx_exporter import export_docx
from sample_solution import fixture


class FixtureLLM:
    def __init__(self, **kwargs):
        self.request, self.blueprint, self.bodies, _ = fixture()
        self.calls = []

    def generate(self, **kwargs):
        key = kwargs["step_key"]
        self.calls.append(key)
        kwargs["state_store"].increment_llm_calls()
        if key == "blueprint":
            assert "visualization_goals" in kwargs["user_prompt"]
            return json.dumps(self.blueprint, ensure_ascii=False)
        if key.startswith("chapter_"):
            assert "```visualizations" in kwargs["user_prompt"]
            return self.bodies[key[len("chapter_"):]]
        raise AssertionError(f"Unexpected extra LLM call: {key}")


def run_fixture(out_dir, failed_render=False):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    req, _, _, _ = fixture()
    path = out_dir/"request.json"
    path.write_text(json.dumps(req, ensure_ascii=False, indent=2), encoding="utf-8")
    args = build_arg_parser().parse_args(["--input-json", str(path), "--output-dir", str(out_dir),
        "--model", "offline-fixture", "--run-mode", "fast", "--research-mode", "off"])
    llm = FixtureLLM()
    with patch("orchestrate_solution.OpenAILLM", return_value=llm):
        if failed_render:
            with patch("solution_skill.visualization.pipeline.render", side_effect=RuntimeError("deliberate render failure")):
                assert orchestrate(args) == 0
        else:
            assert orchestrate(args) == 0
    return llm


@pytest.mark.parametrize("failed_render", [False, True])
def test_actual_orchestrator_and_export(tmp_path, failed_render):
    llm = run_fixture(tmp_path, failed_render)
    state = json.loads((tmp_path/"run_state.json").read_text(encoding="utf-8"))
    assert len(llm.calls) == state["llm_call_count"] == 7  # blueprint + six chapters
    assert state["status"] == "completed" and state["docx_status"] == "completed"
    md = (tmp_path/"solution.md").read_text(encoding="utf-8")
    assert md.count("![") == (0 if failed_render else 5)
    assert "三层服务体系" in md and "```visualizations" not in md
    doc = Document(tmp_path/"sample_solution.docx")
    assert sum(p.text.startswith("图") for p in doc.paragraphs) == (0 if failed_render else 5)
    # Exporter and standalone renderer are the same code, including multiple spacing.
    for table in doc.tables:
        if len(table.rows) > 1:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        assert p.paragraph_format.line_spacing == 1.2
                        assert p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert state["diagram_status"] == ("partial" if failed_render else "completed")


def test_optional_stage_exception_keeps_export(tmp_path):
    with patch("orchestrate_solution.visualize_solution", side_effect=RuntimeError("stage unavailable")):
        llm = run_fixture(tmp_path)
    assert len(llm.calls) == 7 and (tmp_path/"sample_solution.docx").exists()
    state = json.loads((tmp_path/"run_state.json").read_text(encoding="utf-8"))
    assert state["diagram_status"] == "failed" and state["docx_status"] == "completed"
